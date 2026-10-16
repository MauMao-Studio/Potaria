import os
import sys
import shutil

# Ensure project root is in path
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from src.ui.inventory import ITEM_REGISTRY
from src.ui.crafting import RECIPES

# docx imports
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# reportlab imports
import reportlab
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable, Image as RLImage
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

filipino_dishes = [
    'chicken_adobo', 'pork_adobo', 'sinigang', 'kare_kare', 'bulalo', 'tinola',
    'lechon', 'pancit_bihon', 'arroz_caldo', 'halo_halo', 'kaldereta', 'balut',
    'itlog_na_maalat', 'sinangag', 'daing_na_bangus', 'sinigang_na_bangus',
    'ginataang_tilapia', 'kinilaw_na_tanigue'
]

cooked_fishes = [
    'cooked_bangus', 'cooked_tilapia', 'cooked_galunggong', 'cooked_maya_maya',
    'cooked_tanigue', 'cooked_tambakol', 'cooked_lapu_lapu', 'cooked_hito',
    'cooked_dilis', 'cooked_dalagang_bukid'
]

potato_delicacies = [
    'mashed_potato', 'golden_fry', 'hearty_root_stew', 'glazed_golden_tuber',
    'smoked_spud_jerky', 'gilded_russet', 'potato_bread', 'sweet_tuber_pie',
    'herbal_salve'
]

raw_fishes = [
    'raw_bangus', 'raw_tilapia', 'raw_galunggong', 'raw_maya_maya', 'raw_tanigue',
    'raw_tambakol', 'raw_lapu_lapu', 'raw_hito', 'raw_dilis', 'raw_dalagang_bukid'
]

raw_meats = [
    'raw_chicken', 'raw_pork', 'pork_belly', 'raw_beef', 'raw_goat_meat',
    'duck_egg', 'carabao_milk'
]

pantry_ingredients = [
    'raw_potato', 'rice', 'garlic', 'onion', 'ginger', 'tamarind', 'eggplant',
    'string_beans', 'black_pepper', 'calamansi', 'soy_sauce', 'vinegar',
    'salt', 'shaved_ice', 'condensed_milk', 'ube_halaya', 'pancit_noodles',
    'peanut_butter', 'coconut_milk', 'spud_sap', 'seaweed', 'snow_block',
    'wood', 'glowing_mushroom', 'gold_bar'
]

recipe_map = {r["result"]: r for r in RECIPES}

ITEM_ACQUISITION = {
    'raw_potato': 'Harvested from tilled soil or wild potato plants.',
    'rice': 'Harvested from flooded paddy fields or purchased from traders.',
    'garlic': 'Foraged in grassy meadows or grown in garden plots.',
    'onion': 'Grown from seeds or found in crates.',
    'ginger': 'Harvested from root patches along riverbanks.',
    'tamarind': 'Gathered from wild tamarind trees.',
    'eggplant': 'Harvested from farm soil or wild plants.',
    'string_beans': 'Harvested from crop trellises or wild bushes.',
    'black_pepper': 'Harvested from pepper vines in forests.',
    'calamansi': 'Picked from native citrus shrubs near coasts.',
    'soy_sauce': 'Purchased from goods traders or crafted from fermented soybeans.',
    'vinegar': 'Fermented from palm or cane sap in glass jars.',
    'salt': 'Smelted from evaporated sea water or seaweed in a furnace.',
    'shaved_ice': 'Crushed by hand from snow blocks.',
    'condensed_milk': 'Simmered from carabao milk and sugar in a furnace.',
    'ube_halaya': 'Cooked in a furnace with purple yam, milk, and sugar.',
    'pancit_noodles': 'Rolled and cut from rice flour dough.',
    'peanut_butter': 'Roasted and ground peanuts with salt.',
    'coconut_milk': 'Pressed from fresh grated coconut meat.',
    'spud_sap': 'Extracted by tapping mature potato roots.',
    'seaweed': 'Gathered from shallow coastal waters.',
    'snow_block': 'Quarried from snowy mountain peaks.',
    'wood': 'Chopped from trees using an axe.',
    'glowing_mushroom': 'Harvested from cavern biomes.',
    'gold_bar': 'Smelted in a furnace from raw gold ore.',
    'raw_bangus': 'Caught in brackish estuaries with a fishing rod.',
    'raw_tilapia': 'Caught in freshwater rivers and lakes with a fishing rod.',
    'raw_galunggong': 'Caught in coastal ocean waters with a fishing rod.',
    'raw_maya_maya': 'Caught near deep coral reefs with a fishing rod.',
    'raw_tanigue': 'Caught in open ocean waters with a fishing rod.',
    'raw_tambakol': 'Caught in deep ocean waters with a fishing rod.',
    'raw_lapu_lapu': 'Caught around rocky reefs and sea caves with a fishing rod.',
    'raw_hito': 'Caught in muddy freshwater ponds and riverbeds with a fishing rod.',
    'raw_dilis': 'Caught in shallow coastal waters with light tackle or nets.',
    'raw_dalagang_bukid': 'Caught around coral reefs with a fishing rod.',
    'raw_chicken': 'Obtained from poultry in meadows and farms.',
    'raw_pork': 'Obtained from pigs and wild boars.',
    'pork_belly': 'Cut of pork belly from pigs or livestock traders.',
    'raw_beef': 'Obtained from cattle in pasturelands.',
    'raw_goat_meat': 'Obtained from mountain goats.',
    'duck_egg': 'Gathered from duck nests and farm coops.',
    'carabao_milk': 'Milked from water buffalo.'
}

SECTIONS = [
    ("Authentic Filipino Cuisine", filipino_dishes, "Traditional Filipino dishes that restore health and grant specialized meal buffs."),
    ("Fried, Smoked, and Steamed Fish", cooked_fishes, "Cooked freshwater and marine fish that restore health and provide combat stat buffs."),
    ("Potato Specialities and Delicacies", potato_delicacies, "Potato dishes, breads, and restorative salves crafted from farm tubers."),
    ("Fresh Catches (Raw Fishes)", raw_fishes, "Fresh raw fish caught using fishing rods across freshwater and saltwater biomes."),
    ("Livestock, Poultry, and Dairy", raw_meats, "Raw meats, poultry, eggs, and fresh milk gathered from domestic animals."),
    ("Pantry, Spices, and Cooking Essentials", pantry_ingredients, "Grains, seasonings, and basic ingredients used for cooking recipes.")
]

def get_item_icon_path(item_id):
    """
    Resolves the sprite path for an item, looking in assets/items and assets/craftable_items.
    """
    candidates = [
        os.path.join(BASE_DIR, "assets", "items", f"{item_id}.png"),
        os.path.join(BASE_DIR, "assets", "craftable_items", f"{item_id}.png"),
        os.path.join("assets", "items", f"{item_id}.png"),
        os.path.join("assets", "craftable_items", f"{item_id}.png"),
    ]
    for p in candidates:
        if os.path.isfile(p):
            return p
    return None

def format_health_and_buff(data):
    """Formats the health column to include health restore and meal buff category/effects."""
    heal_val = data.get("heal", 0)
    heal_str = f"+{heal_val} HP" if heal_val > 0 else "Base"

    buff = data.get("buff")
    if buff:
        b_cat = str(buff.get("category", "")).upper()
        dur_s = int(buff.get("duration", 180))
        dur_str = f"{dur_s // 60}m" if dur_s >= 60 else f"{dur_s}s"
        eff_list = data.get("effects", [])
        if eff_list:
            eff_text = ", ".join(eff_list)
            heal_str += f"\n[{b_cat}] {eff_text}"
        elif buff.get("stats"):
            stat_parts = []
            for sk, sv in buff["stats"].items():
                s_name = sk.replace("_mult", "").replace("_", " ").title()
                s_val = f"+{sv*100:.0f}%" if "mult" in sk or "chance" in sk or "pen" in sk or "reduction" in sk else f"+{sv}"
                stat_parts.append(f"{s_val} {s_name}")
            eff_text = ", ".join(stat_parts)
            heal_str += f"\n[{b_cat}] {eff_text} ({dur_str})"
    return heal_str, heal_val

# -------------------------------------------------------------
# DOCX BUILDER
# -------------------------------------------------------------
def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def build_docx(with_codes=True, output_filename=None):
    if output_filename is None:
        output_filename = os.path.join(
            BASE_DIR,
            "Potaria_Cookbook_With_Codes.docx" if with_codes else "Potaria_Cookbook_No_Codes.docx"
        )

    doc = Document()
    
    # Page setup: 0.7 in margins -> 7.1 in printable width
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.7)
        s.right_margin = Inches(0.7)
        s.different_first_page_header_footer = True
        
        # Header on later pages
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("THE ROYAL POTARIA COOKBOOK | Culinary Compendium")
        hrun.font.name = "Georgia"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(140, 100, 40)
        
        # Footer
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("The Royal Culinary Guild of Potaria - Official Culinary Records")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(120, 120, 120)

    # Title Cover
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("THE ROYAL POTARIA COOKBOOK")
    r_title.bold = True
    r_title.font.name = "Georgia"
    r_title.font.size = Pt(26)
    r_title.font.color.rgb = RGBColor(123, 17, 19) # Burgundy
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_suffix = " (Reference Edition with Item Codes)" if with_codes else " (Adventurer Edition)"
    r_sub = p_sub.add_run(f"An Epicurean Guide to Culinary Arts, Hearth Recipes, and Native Ingredients of Potaria{subtitle_suffix}")
    r_sub.italic = True
    r_sub.font.name = "Georgia"
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(140, 100, 40) # Warm bronze
    
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_desc = "Published by The Royal Culinary Guild of Potaria | Complete Edition 1.0 (Developer & Crafting Code Index)" if with_codes else "Published by The Royal Culinary Guild of Potaria | Complete Edition 1.0 (Player Immersion Edition)"
    r_meta = p_meta.add_run(f"{meta_desc}\n")
    r_meta.font.name = "Calibri"
    r_meta.font.size = Pt(10)
    r_meta.font.color.rgb = RGBColor(100, 100, 100)
    
    doc.add_paragraph() # spacer
    
    # Introduction Section
    p_h1 = doc.add_heading(level=1)
    r_h1 = p_h1.add_run("1. Principles of Hearth and Culinary Craft")
    r_h1.font.name = "Georgia"
    r_h1.font.size = Pt(16)
    r_h1.font.color.rgb = RGBColor(123, 17, 19)
    
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.line_spacing = 1.15
    p_intro.paragraph_format.space_after = Pt(8)
    p_intro.add_run(
        "Culinary craft in the world of Potaria represents one of the most vital disciplines for survival, exploration, "
        "and legendary triumph. While potions offer instantaneous medicinal alchemy, prepared meals provide profound vitality, "
        "replenishing substantial Health Points (HP) and conferring tactical nourishment through the 5-Active-Meal-Buff Architecture.\n\n"
        "Under this gastronomic discipline, adventurers may maintain up to five distinct active food buffs simultaneously—strictly "
        "one meal per category: Offense, Defense, Mobility, Resource, and Utility. Consuming a subsequent dish within an already "
        "occupied category cleanly replaces the existing buff with the new recipe's full authored duration, preventing runaway magnitude "
        "or duration stacking while empowering deliberate, strategic feast crafting."
    )
    
    p_stations = doc.add_paragraph()
    p_stations.paragraph_format.line_spacing = 1.15
    p_stations.paragraph_format.space_after = Pt(12)
    p_stations.add_run("Potaria features two recognized modes of culinary assembly:\n").bold = True
    p_stations.add_run("• Furnace Hearth Cooking (station=\"furnace\"): ").bold = True
    p_stations.add_run("The traditional hearth utilizes fuel combustion to simmer complex stews (Bulalo, Kare-Kare, Sinigang), deep-fry marine catches, roast prime meats (Lechon Kawali), bake breads, and clarify condensed milks.\n")
    p_stations.add_run("• Artisanal Preparation By Hand (station=None): ").bold = True
    p_stations.add_run("Delicate refreshments, iced desserts (Special Halo-Halo), cured marinades (Kinilaw na Tanigue), fresh brined duck eggs, and noodle rolling require no thermal heat, executed directly from the player's personal craft inventory.")
    
    # Column configuration (Total = 7.10 inches)
    if with_codes:
        # 7 columns
        col_widths = [Inches(0.45), Inches(1.30), Inches(0.95), Inches(0.85), Inches(0.70), Inches(1.65), Inches(1.20)]
        hdr_titles = ["Icon", "Dish / Item Name", "Item ID", "Health", "Station", "Ingredients / Source", "Culinary Notes"]
    else:
        # 6 columns
        col_widths = [Inches(0.45), Inches(1.45), Inches(0.95), Inches(0.75), Inches(2.10), Inches(1.40)]
        hdr_titles = ["Icon", "Dish / Item Name", "Health", "Station", "Ingredients / Source", "Culinary Notes"]
    
    for sec_idx, (title, items, desc_text) in enumerate(SECTIONS, start=2):
        doc.add_page_break()
        
        h_sec = doc.add_heading(level=1)
        r_sec = h_sec.add_run(f"{sec_idx}. {title}")
        r_sec.font.name = "Georgia"
        r_sec.font.size = Pt(15)
        r_sec.font.color.rgb = RGBColor(123, 17, 19)
        
        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.line_spacing = 1.15
        p_desc.paragraph_format.space_after = Pt(10)
        p_desc.add_run(desc_text).italic = True
        
        # Create Table
        table = doc.add_table(rows=1, cols=len(hdr_titles))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        
        # Repeat header row across pages
        trPr = table.rows[0]._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        
        # Header Row
        hdr_cells = table.rows[0].cells
        for i, htitle in enumerate(hdr_titles):
            hdr_cells[i].width = col_widths[i]
            set_cell_background(hdr_cells[i], "7B1113")
            set_cell_margins(hdr_cells[i], top=110, bottom=110, left=40 if i == 0 else 80, right=40 if i == 0 else 80)
            hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            hdr_p = hdr_cells[i].paragraphs[0]
            hdr_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            hdr_run = hdr_p.add_run(htitle)
            hdr_run.bold = True
            hdr_run.font.name = "Calibri"
            hdr_run.font.size = Pt(9.5)
            hdr_run.font.color.rgb = RGBColor(255, 255, 255)
        
        # Data Rows
        for r_idx, item_id in enumerate(items):
            row = table.add_row()
            # CantSplit property prevents awkward row breaking across pages
            r_trPr = row._tr.get_or_add_trPr()
            r_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            
            row_cells = row.cells
            bg_color = "FAF6F0" if r_idx % 2 == 1 else "FFFFFF"
            
            data = ITEM_REGISTRY[item_id]
            name = data.get("name", item_id)
            heal_str, heal_val = format_health_and_buff(data)
            desc_val = data.get("desc", "")
            
            rec = recipe_map.get(item_id)
            if rec:
                st_str = "Furnace" if rec.get("station") == "furnace" else "By Hand"
                req_list = [f"{cnt}x {ITEM_REGISTRY.get(iid, {}).get('name', iid)}" for iid, cnt in rec.get("ingredients", [])]
                ing_str = ", ".join(req_list)
            else:
                st_str = "Forage/Hunt"
                ing_str = ITEM_ACQUISITION.get(item_id, "Natural harvest")
            
            icon_path = get_item_icon_path(item_id)
            
            # Format row data according to version
            if with_codes:
                text_columns = [
                    (1, name, WD_ALIGN_PARAGRAPH.LEFT, True, Pt(8.5), RGBColor(40, 40, 40)),
                    (2, item_id, WD_ALIGN_PARAGRAPH.CENTER, False, Pt(7.5), RGBColor(100, 100, 100)),
                    (3, heal_str, WD_ALIGN_PARAGRAPH.CENTER, heal_val > 0, Pt(7.5), RGBColor(34, 139, 34) if heal_val > 0 else RGBColor(50, 50, 50)),
                    (4, st_str, WD_ALIGN_PARAGRAPH.CENTER, False, Pt(8.5), RGBColor(50, 50, 50)),
                    (5, ing_str, WD_ALIGN_PARAGRAPH.LEFT, False, Pt(8.5), RGBColor(50, 50, 50)),
                    (6, desc_val, WD_ALIGN_PARAGRAPH.LEFT, False, Pt(8.0), RGBColor(80, 80, 80)),
                ]
            else:
                text_columns = [
                    (1, name, WD_ALIGN_PARAGRAPH.LEFT, True, Pt(8.5), RGBColor(40, 40, 40)),
                    (2, heal_str, WD_ALIGN_PARAGRAPH.CENTER, heal_val > 0, Pt(7.5), RGBColor(34, 139, 34) if heal_val > 0 else RGBColor(50, 50, 50)),
                    (3, st_str, WD_ALIGN_PARAGRAPH.CENTER, False, Pt(8.5), RGBColor(50, 50, 50)),
                    (4, ing_str, WD_ALIGN_PARAGRAPH.LEFT, False, Pt(8.5), RGBColor(50, 50, 50)),
                    (5, desc_val, WD_ALIGN_PARAGRAPH.LEFT, False, Pt(8.0), RGBColor(80, 80, 80)),
                ]
            
            # Set up Icon cell (Col 0)
            icon_cell = row_cells[0]
            icon_cell.width = col_widths[0]
            set_cell_background(icon_cell, bg_color)
            set_cell_margins(icon_cell, top=60, bottom=60, left=40, right=40)
            icon_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            icon_p = icon_cell.paragraphs[0]
            icon_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            icon_p.paragraph_format.space_before = Pt(0)
            icon_p.paragraph_format.space_after = Pt(0)
            if icon_path and os.path.exists(icon_path):
                run = icon_p.add_run()
                run.add_picture(icon_path, width=Inches(0.32))
            
            # Set up text columns
            for c_i, val, alignment, is_bold, fsize, fcolor in text_columns:
                cell = row_cells[c_i]
                cell.width = col_widths[c_i]
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=80, bottom=80, left=70, right=70)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                cell_p = cell.paragraphs[0]
                cell_p.alignment = alignment
                cell_p.paragraph_format.line_spacing = 1.05
                cell_p.paragraph_format.space_before = Pt(0)
                cell_p.paragraph_format.space_after = Pt(2)
                
                run = cell_p.add_run(val)
                run.font.name = "Consolas" if (with_codes and c_i == 2) else "Calibri"
                run.font.size = fsize
                run.bold = is_bold
                run.font.color.rgb = fcolor
                if not with_codes and c_i == 5 or with_codes and c_i == 6:
                    run.italic = True
                    
        # Apply table column width definitions
        for idx, width in enumerate(col_widths):
            table.columns[idx].width = width
    
    # Save Word Doc
    doc.save(output_filename)
    print(f"Generated Word Document: {output_filename}")
    return output_filename

# -------------------------------------------------------------
# REPORTLAB PDF BUILDER
# -------------------------------------------------------------
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super().showPage()
        super().save()

    def draw_page_decorations(self, total_pages):
        page_num = self._pageNumber
        self.saveState()
        
        # Header (Pages > 1)
        if page_num > 1:
            self.setFont("Helvetica-Bold", 8)
            self.setFillColor(colors.HexColor("#7B1113"))
            self.drawString(36, 755, "THE ROYAL POTARIA COOKBOOK")
            self.setFont("Helvetica-Oblique", 8)
            self.setFillColor(colors.HexColor("#777777"))
            self.drawRightString(612 - 36, 755, "Culinary Compendium & Hearth Guide")
            self.setStrokeColor(colors.HexColor("#D4AF37"))
            self.setLineWidth(0.5)
            self.line(36, 750, 612 - 36, 750)
            
        # Footer
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#777777"))
        self.drawRightString(612 - 36, 32, f"Page {page_num} of {total_pages}")
        self.drawString(36, 32, "Potaria Official Culinary Records")
        self.restoreState()

def draw_page_background(canvas_obj, doc_obj):
    canvas_obj.saveState()
    # Draw background tint BEFORE story elements
    canvas_obj.setFillColor(colors.HexColor("#FAF7F2"))
    canvas_obj.rect(0, 0, 612, 792, stroke=0, fill=1)
    
    # Decorative border
    canvas_obj.setStrokeColor(colors.HexColor("#D4AF37")) # Gold
    canvas_obj.setLineWidth(1.5)
    canvas_obj.rect(24, 24, 612 - 48, 792 - 48, stroke=1, fill=0)
    canvas_obj.setStrokeColor(colors.HexColor("#7B1113")) # Inner Burgundy
    canvas_obj.setLineWidth(0.5)
    canvas_obj.rect(27, 27, 612 - 54, 792 - 54, stroke=1, fill=0)
    canvas_obj.restoreState()

def build_pdf(with_codes=True, output_filename=None):
    if output_filename is None:
        output_filename = os.path.join(
            BASE_DIR,
            "Potaria_Cookbook_With_Codes.pdf" if with_codes else "Potaria_Cookbook_No_Codes.pdf"
        )
    
    # 36pt margins, page width 612pt -> printable width 540pt
    doc = SimpleDocTemplate(
        output_filename,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=46,
        bottomMargin=46
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=24,
        leading=28,
        textColor=colors.HexColor('#7B1113'),
        alignment=1, # Center
        spaceAfter=8
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#8B5A2B'),
        alignment=1,
        spaceAfter=14
    )
    
    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor('#666666'),
        alignment=1,
        spaceAfter=18
    )
    
    h1_style = ParagraphStyle(
        'SectionH1',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=17,
        textColor=colors.HexColor('#7B1113'),
        spaceBefore=10,
        spaceAfter=6
    )
    
    body_style = ParagraphStyle(
        'SectionBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor('#2E2E2E'),
        spaceAfter=8
    )
    
    th_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1
    )
    
    td_name_style = ParagraphStyle(
        'TableDishName',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#1A1A1A')
    )
    
    td_id_style = ParagraphStyle(
        'TableDishID',
        parent=styles['Normal'],
        fontName='Courier',
        fontSize=6.0,
        leading=7.5,
        textColor=colors.HexColor('#555555'),
        alignment=1
    )
    
    td_heal_style = ParagraphStyle(
        'TableHeal',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=6.5,
        leading=8.0,
        textColor=colors.HexColor('#1E7E34'),
        alignment=1
    )
    
    td_heal_base_style = ParagraphStyle(
        'TableHealBase',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=6.5,
        leading=8.0,
        textColor=colors.HexColor('#555555'),
        alignment=1
    )
    
    td_st_style = ParagraphStyle(
        'TableStation',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.0,
        leading=8.5,
        textColor=colors.HexColor('#4A4A4A'),
        alignment=1
    )
    
    td_ing_style = ParagraphStyle(
        'TableIngredients',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor('#222222')
    )
    
    td_lore_style = ParagraphStyle(
        'TableLore',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=6.5,
        leading=8.5,
        textColor=colors.HexColor('#444444')
    )
    
    story = []
    
    # Title & Front Matter
    story.append(Spacer(1, 15))
    story.append(Paragraph("THE ROYAL POTARIA COOKBOOK", title_style))
    subtitle_suffix = " (Reference Edition with Item Codes)" if with_codes else " (Adventurer Edition)"
    story.append(Paragraph(f"A Comprehensive Culinary Compendium & Epicurean Guide to the Wilds of Potaria{subtitle_suffix}", subtitle_style))
    meta_desc = "Authored by the Master Gastronomers of the Realm | Reference Edition with Item Codes (Release 1.0.0)" if with_codes else "Authored by the Master Gastronomers of the Realm | Standard Adventurer Edition (Release 1.0.0)"
    story.append(Paragraph(meta_desc, meta_style))
    story.append(HRFlowable(width="90%", thickness=1.5, color=colors.HexColor("#D4AF37"), spaceAfter=15, spaceBefore=0))
    
    story.append(Paragraph("Culinary Framework & Gastronomic Lore", h1_style))
    story.append(Paragraph(
        "Cooking in Potaria provides vital sustenance, restorative healing, and tactical nourishment through the 5-Active-Meal-Buff Architecture. "
        "Adventurers may sustain up to five simultaneous meal buffs—strictly one per discipline: Offense, Defense, Mobility, Resource, and Utility. "
        "Eating a dish within an active category replaces its predecessor with fresh duration, ensuring balanced tactical customization. "
        "Recipes are prepared across two core traditions: "
        "<b>Furnace Hearth Cooking</b> (requiring thermal combustion to simmer broths, fry catches, and bake loaves) and "
        "<b>By-Hand Preparation</b> (delicate cold preparations, cured ceviches like Kinilaw na Tanigue, iced halo-halo, and herbal salves).",
        body_style
    ))
    
    story.append(Spacer(1, 10))
    
    # 540 pt total width
    if with_codes:
        # 7 columns: 30 + 90 + 78 + 60 + 50 + 127 + 105 = 540 pt
        col_widths = [30, 90, 78, 60, 50, 127, 105]
        hdr_cells = [
            Paragraph("Icon", th_style),
            Paragraph("Dish / Item Name", th_style),
            Paragraph("Item ID", th_style),
            Paragraph("Health", th_style),
            Paragraph("Station", th_style),
            Paragraph("Recipe Ingredients / Source", th_style),
            Paragraph("Culinary Notes", th_style)
        ]
    else:
        # 6 columns: 32 + 110 + 66 + 52 + 155 + 125 = 540 pt
        col_widths = [32, 110, 66, 52, 155, 125]
        hdr_cells = [
            Paragraph("Icon", th_style),
            Paragraph("Dish / Item Name", th_style),
            Paragraph("Health", th_style),
            Paragraph("Station", th_style),
            Paragraph("Recipe Ingredients / Source", th_style),
            Paragraph("Culinary Notes", th_style)
        ]
    
    for sec_idx, (title, items, desc_text) in enumerate(SECTIONS, start=1):
        if sec_idx > 1:
            story.append(PageBreak())
        
        story.append(Paragraph(f"Section {sec_idx}: {title}", h1_style))
        story.append(Paragraph(desc_text, body_style))
        story.append(Spacer(1, 4))
        
        table_data = [hdr_cells]
        
        for r_idx, item_id in enumerate(items):
            data = ITEM_REGISTRY[item_id]
            name = data.get("name", item_id)
            heal_str, heal_val = format_health_and_buff(data)
            desc_val = data.get("desc", "")
            
            rec = recipe_map.get(item_id)
            if rec:
                st_str = "Furnace" if rec.get("station") == "furnace" else "By Hand"
                req_list = [f"{cnt}x {ITEM_REGISTRY.get(iid, {}).get('name', iid)}" for iid, cnt in rec.get("ingredients", [])]
                ing_str = ", ".join(req_list)
            else:
                st_str = "Forage/Hunt"
                ing_str = ITEM_ACQUISITION.get(item_id, "Natural harvest")
            
            icon_path = get_item_icon_path(item_id)
            if icon_path and os.path.exists(icon_path):
                icon_flowable = RLImage(icon_path, width=22, height=22)
                icon_flowable.hAlign = 'CENTER'
            else:
                icon_flowable = ""
            
            p_name = Paragraph(name, td_name_style)
            p_heal = Paragraph(heal_str.replace("\n", "<br/>"), td_heal_style if heal_val > 0 else td_heal_base_style)
            p_st = Paragraph(st_str, td_st_style)
            p_ing = Paragraph(ing_str, td_ing_style)
            p_lore = Paragraph(desc_val, td_lore_style)
            
            if with_codes:
                p_id = Paragraph(item_id, td_id_style)
                table_data.append([icon_flowable, p_name, p_id, p_heal, p_st, p_ing, p_lore])
            else:
                table_data.append([icon_flowable, p_name, p_heal, p_st, p_ing, p_lore])
            
        t = Table(table_data, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#7B1113')),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#DECBB5')),
            ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#7B1113')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5EFEB')]),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('LEFTPADDING', (0, 0), (-1, -1), 3.5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 3.5),
            ('LEFTPADDING', (0, 0), (0, -1), 1.5),
            ('RIGHTPADDING', (0, 0), (0, -1), 1.5),
            ('LEFTPADDING', (3, 0), (4, -1), 2),
            ('RIGHTPADDING', (3, 0), (4, -1), 2),
        ]))
        
        story.append(t)
        story.append(Spacer(1, 10))
        
    doc.build(story, onFirstPage=draw_page_background, onLaterPages=draw_page_background, canvasmaker=NumberedCanvas)
    print(f"Generated PDF Document: {output_filename}")
    return output_filename

def build_all(output_dir=None):
    """Generates all 4 cookbook files and synchronizes primary aliases."""
    if output_dir is None:
        output_dir = BASE_DIR
    os.makedirs(output_dir, exist_ok=True)
    print("Building Potaria Cookbook Documents...")
    
    # 1. Version 1: With Codes
    docx_codes_path = os.path.join(output_dir, "Potaria_Cookbook_With_Codes.docx")
    docx_codes = build_docx(with_codes=True, output_filename=docx_codes_path)
    shutil.copyfile(docx_codes, os.path.join(output_dir, "Potaria_Cookbook.docx"))
    print("Synchronized Potaria_Cookbook.docx -> Potaria_Cookbook_With_Codes.docx")
    
    pdf_codes_path = os.path.join(output_dir, "Potaria_Cookbook_With_Codes.pdf")
    pdf_codes = build_pdf(with_codes=True, output_filename=pdf_codes_path)
    shutil.copyfile(pdf_codes, os.path.join(output_dir, "Potaria_Cookbook.pdf"))
    print("Synchronized Potaria_Cookbook.pdf -> Potaria_Cookbook_With_Codes.pdf")
    
    # 2. Version 2: Without Codes
    docx_no_codes_path = os.path.join(output_dir, "Potaria_Cookbook_No_Codes.docx")
    build_docx(with_codes=False, output_filename=docx_no_codes_path)
    
    pdf_no_codes_path = os.path.join(output_dir, "Potaria_Cookbook_No_Codes.pdf")
    build_pdf(with_codes=False, output_filename=pdf_no_codes_path)
    
    print("All 4 Cookbook documents generated successfully.")

if __name__ == "__main__":
    build_all()
