# -*- coding: utf-8 -*-

import sys
import os
import io
import json

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)

import pygame
pygame.init()
pygame.display.set_mode((1, 1), pygame.NOFRAME)

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.shared import OxmlElement
from docx.oxml.ns import qn

from src.assets.asset_factory import AssetFactory
from src.ui.inventory import ITEM_REGISTRY
import src.ui.crafting  # Ensures all 129 craftable items are registered in ITEM_REGISTRY

CATEGORY_NAMES = {
    "tool": "Tools & Mining Equipment",
    "weapon": "Weapons (Melee, Ranged, Magic)",
    "armor_head": "Armor - Helmets & Headwear",
    "armor_chest": "Armor - Chestplates & Robes",
    "armor_legs": "Armor - Greaves & Boots",
    "accessory": "Accessories, Rings & Talismans",
    "skillbook": "Skill Books & Spell Tomes",
    "block": "Blocks & Placeable Tiles",
    "wall": "Background Walls",
    "consumable": "Consumables & Food",
    "special": "Special & Quest Relics",
    "material": "Crafting Materials & Drops",
    "currency": "Currency & Coins",
}

RARITY_COLORS_DOCX = {
    "common": RGBColor(127, 140, 141),
    "uncommon": RGBColor(39, 174, 96),
    "rare": RGBColor(41, 128, 185),
    "epic": RGBColor(142, 68, 173),
    "mythic": RGBColor(211, 84, 0),
    "legendary": RGBColor(212, 172, 13),
    "ancient": RGBColor(192, 57, 43),
}

RARITY_ORDER = ["common", "uncommon", "rare", "epic", "mythic", "legendary", "ancient"]


def set_repeat_table_header(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cantSplit = OxmlElement("w:cantSplit")
    cantSplit.set(qn("w:val"), "true")
    trPr.append(cantSplit)


def shade_header_cells(hdr_cells, headers):
    for i, text in enumerate(headers):
        hdr_cells[i].text = text
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(255, 255, 255)
        tc = hdr_cells[i]._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "2C3440")
        tcPr.append(shd)


def set_col_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width


def generate_docx():
    af = AssetFactory()
    document = Document()

    # Landscape Page Setup with dedicated header/footer margins
    section = document.sections[0]
    new_width, new_height = section.page_height, section.page_width
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = new_width
    section.page_height = new_height
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    total_items = len(ITEM_REGISTRY)

    header = section.header
    header_p = header.paragraphs[0]
    header_p.text = "POTARIA: ROOT REALM - Official Visual Item & Affix Catalog"
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.text = f"Developer: Kurados | Total Registered Items: {total_items} | Potaria & Potaria: Root Realm"
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Document Title Block
    title = document.add_heading("POTARIA: ROOT REALM", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Official Visual Item Catalog, Crafting Reference & Data-Driven Affix Guide")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info = document.add_paragraph(
        f"Total Registered Items: {total_items} (All Craftable & Droppable Items) | Developer: Kurados"
    )
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Section 1: Data-Driven Affix Registry Table
    affixes_path = os.path.join(BASE_DIR, "data", "affixes", "affixes.json")
    if os.path.exists(affixes_path):
        with open(affixes_path, "r", encoding="utf-8") as f:
            affixes_data = json.load(f)

        document.add_heading(f"Data-Driven Multi-Affix Registry ({len(affixes_data)} Modifiers)", level=1)
        aff_table = document.add_table(rows=1, cols=6)
        aff_table.style = "Table Grid"
        shade_header_cells(
            aff_table.rows[0].cells,
            ["Affix Key", "Slot Type", "Power Type", "Min Rarity", "Eligible Categories", "Effect Description"],
        )
        set_repeat_table_header(aff_table.rows[0])
        prevent_row_split(aff_table.rows[0])

        for aff_key, aff in affixes_data.items():
            row = aff_table.add_row()
            prevent_row_split(row)
            cells = row.cells
            cells[0].text = aff.get("name", aff_key)
            cells[1].text = aff.get("type", "prefix").title()
            cells[2].text = aff.get("power_type", "stat").replace("_", " ").title()
            cells[3].text = aff.get("min_rarity", "Uncommon")
            cells[4].text = ", ".join(aff.get("allowed_cats", []))
            cells[5].text = aff.get("desc", "")

        set_col_widths(
            aff_table,
            [Inches(1.3), Inches(0.9), Inches(1.2), Inches(1.0), Inches(2.2), Inches(3.4)],
        )

    # Group items by category
    items_by_cat = {}
    for item_id, item_data in ITEM_REGISTRY.items():
        cat = item_data.get("cat", "other")
        items_by_cat.setdefault(cat, []).append((item_id, item_data))

    cat_sequence = [
        "tool", "weapon", "armor_head", "armor_chest", "armor_legs",
        "accessory", "skillbook", "consumable", "material", "block",
        "wall", "special", "currency",
    ]
    for c in items_by_cat.keys():
        if c not in cat_sequence:
            cat_sequence.append(c)

    item_col_widths = [Inches(0.55), Inches(1.65), Inches(1.75), Inches(0.95), Inches(1.85), Inches(3.25)]

    for cat in cat_sequence:
        if cat not in items_by_cat:
            continue

        items_list = items_by_cat[cat]

        def sort_key(item_tuple):
            iid, d = item_tuple
            rar = str(d.get("rarity", "common")).lower()
            rar_idx = RARITY_ORDER.index(rar) if rar in RARITY_ORDER else 99
            name = d.get("name", iid).lower()
            return (rar_idx, name)

        items_list.sort(key=sort_key)

        cat_display = CATEGORY_NAMES.get(cat, cat.replace("_", " ").title())
        heading = document.add_heading(f"{cat_display} ({len(items_list)} items)", level=1)
        heading.paragraph_format.keep_with_next = True

        table = document.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        shade_header_cells(
            table.rows[0].cells,
            ["Icon", "Item Code", "Item Name", "Rarity", "Stats", "Description"],
        )
        set_repeat_table_header(table.rows[0])
        prevent_row_split(table.rows[0])

        for item_id, d in items_list:
            row = table.add_row()
            prevent_row_split(row)
            row_cells = row.cells

            icon_surf = af.items.get(item_id)
            if not icon_surf:
                icon_surf = af.generate_item_icon(item_id)
            upscaled_surf = pygame.transform.scale(icon_surf, (56, 56))

            buf = io.BytesIO()
            pygame.image.save(upscaled_surf, buf, "png")
            buf.seek(0)

            p = row_cells[0].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run()
            r.add_picture(buf, width=Inches(0.32))

            row_cells[1].text = item_id
            row_cells[2].text = d.get("name", item_id)

            rarity = str(d.get("rarity", "common")).upper()
            row_cells[3].text = ""
            p_rarity = row_cells[3].paragraphs[0]
            run_rarity = p_rarity.add_run(rarity)
            run_rarity.font.bold = True
            run_rarity.font.color.rgb = RARITY_COLORS_DOCX.get(rarity.lower(), RGBColor(0, 0, 0))

            stat_parts = []
            for stat in ["damage", "atk", "def", "crit", "speed", "hp", "mp", "mp_regen", "heal", "special"]:
                if d.get(stat, 0):
                    stat_parts.append(f"{stat.upper()}: {d[stat]}")
            stats_str = ", ".join(stat_parts) if stat_parts else "-"
            row_cells[4].text = stats_str

            desc = d.get("desc", "") or d.get("lore", "") or "-"
            row_cells[5].text = desc

        set_col_widths(table, item_col_widths)

    out_path = os.path.join(BASE_DIR, "Potaria_Items_Catalog.docx")
    document.save(out_path)
    print(f"Generated DOCX at: {out_path} (Items: {total_items})")


if __name__ == "__main__":
    generate_docx()
